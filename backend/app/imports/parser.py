"""Safe CSV and XLSX file parsing and Unicode value normalization."""

import csv
import datetime
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any

import openpyxl  # type: ignore[import-untyped]

from backend.app.core.exceptions import ValidationException


@dataclass
class ParsedRow:
    """Represents a single parsed row from an import file."""

    row_number: int
    raw_data: dict[str, Any]
    parse_error: str | None = None


@dataclass
class ParsedImport:
    """Encapsulates parsed file contents and metadata."""

    file_hash: str
    file_size: int
    headers: list[str]
    raw_headers: list[str]
    rows: list[ParsedRow] = field(default_factory=list)
    sheet_names: list[str] = field(default_factory=list)
    selected_sheet: str | None = None


def normalize_header(header: str) -> str:
    """Normalize a column header name to a standard snake_case key.

    Example: 'Student ID' -> 'student_id', 'First-Name' -> 'first_name'.
    """
    cleaned = unicodedata.normalize("NFKC", str(header or "")).strip()
    cleaned = re.sub(r"[\s\-\.]+", "_", cleaned)
    cleaned = re.sub(r"[^\w_]", "", cleaned)
    return cleaned.lower().strip("_")


def normalize_cell_value(val: Any) -> Any:
    """Normalize cell value: NFKC Unicode, trim whitespace, date/time formatting.

    Preserves Afghan/Perso-Arabic letters (Pashto/Dari) and standard symbols.
    Converts empty or whitespace-only strings to None.
    """
    if val is None:
        return None

    if isinstance(val, bool):
        return val

    if isinstance(val, (datetime.datetime, datetime.date)):
        if isinstance(val, datetime.datetime):
            # If time is 00:00:00 and was entered as date, return date string
            if val.hour == 0 and val.minute == 0 and val.second == 0:
                return val.strftime("%Y-%m-%d")
            return val.isoformat()
        return val.strftime("%Y-%m-%d")

    if isinstance(val, datetime.time):
        return val.strftime("%H:%M:%S")

    if isinstance(val, (int, float)):
        if isinstance(val, float) and val.is_integer():
            return int(val)
        return val

    if isinstance(val, str):
        # Apply NFKC Unicode normalization
        normalized = unicodedata.normalize("NFKC", val)
        # Strip leading and trailing whitespace
        trimmed = normalized.strip()
        if not trimmed:
            return None
        return trimmed

    # Fallback to string
    s_val = unicodedata.normalize("NFKC", str(val)).strip()
    return s_val if s_val else None


def parse_csv_content(
    content: bytes,
    max_rows: int,
) -> tuple[list[str], list[str], list[ParsedRow]]:
    """Parse CSV content with UTF-8 / UTF-8-BOM decoding and header normalization."""
    # Decode content trying utf-8-sig (handles BOM), then utf-8, then latin-1
    text: str
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception as exc:
                raise ValidationException(
                    "Unable to decode CSV file. Supported: UTF-8, UTF-8-BOM, or Latin-1.",
                    details={"error": str(exc)},
                ) from exc

    # Sniff delimiter or default to comma
    sample = text[:4096]
    delimiter = ","
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except Exception:
        delimiter = ","

    stream = io.StringIO(text)
    reader = csv.reader(stream, delimiter=delimiter)

    raw_headers: list[str] = []
    headers: list[str] = []
    rows: list[ParsedRow] = []

    physical_row_num = 0
    header_found = False

    for row in reader:
        physical_row_num += 1
        # Skip completely blank lines
        if not row or all(c.strip() == "" for c in row):
            continue

        if not header_found:
            raw_headers = [str(c).strip() for c in row]
            headers = [normalize_header(c) for c in raw_headers]
            header_found = True
            continue

        if len(rows) >= max_rows:
            raise ValidationException(
                f"File contains more than the maximum permitted {max_rows} data rows.",
                details={"max_rows": max_rows},
            )

        # Map row columns to normalized headers
        row_dict: dict[str, Any] = {}
        for idx, col_name in enumerate(headers):
            if not col_name:
                continue
            raw_val = row[idx] if idx < len(row) else None
            row_dict[col_name] = normalize_cell_value(raw_val)

        rows.append(ParsedRow(row_number=physical_row_num, raw_data=row_dict))

    if not header_found:
        raise ValidationException("CSV file contains no header row or data.")

    return raw_headers, headers, rows


def parse_xlsx_content(
    content: bytes,
    max_rows: int,
    target_sheet: str | None = None,
) -> tuple[list[str], str, list[str], list[str], list[ParsedRow]]:
    """Parse XLSX content safely without formula execution using read_only mode."""
    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise ValidationException(
            "Invalid Excel (.xlsx) spreadsheet format or corrupted file.",
            details={"error": str(exc)},
        ) from exc

    sheet_names = wb.sheetnames
    if not sheet_names:
        wb.close()
        raise ValidationException("Excel workbook contains no worksheets.")

    selected_sheet_name = (
        target_sheet if target_sheet and target_sheet in sheet_names else sheet_names[0]
    )
    ws = wb[selected_sheet_name]

    raw_headers: list[str] = []
    headers: list[str] = []
    rows: list[ParsedRow] = []

    physical_row_num = 0
    header_found = False

    try:
        for row in ws.iter_rows(values_only=True):
            physical_row_num += 1
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue

            if not header_found:
                raw_headers = [str(c).strip() if c is not None else "" for c in row]
                headers = [normalize_header(c) for c in raw_headers]
                header_found = True
                continue

            if len(rows) >= max_rows:
                raise ValidationException(
                    f"Spreadsheet contains more than the maximum permitted {max_rows} data rows.",
                    details={"max_rows": max_rows},
                )

            row_dict: dict[str, Any] = {}
            for idx, col_name in enumerate(headers):
                if not col_name:
                    continue
                raw_val = row[idx] if idx < len(row) else None
                row_dict[col_name] = normalize_cell_value(raw_val)

            rows.append(ParsedRow(row_number=physical_row_num, raw_data=row_dict))
    finally:
        wb.close()

    if not header_found:
        raise ValidationException("Worksheet contains no header row or data.")

    return sheet_names, selected_sheet_name, raw_headers, headers, rows


def parse_import_file(
    filename: str,
    content: bytes,
    max_bytes: int,
    max_rows: int,
    sheet_name: str | None = None,
) -> ParsedImport:
    """Safe dispatcher for CSV and XLSX file parsing.

    Validates:
    - Content length does not exceed max_bytes.
    - Extension must be .csv, .tsv, .txt, or .xlsx.
    - Row count does not exceed max_rows.
    - Computes SHA-256 digest.
    """
    file_size = len(content)
    if file_size == 0:
        raise ValidationException("The uploaded file is empty (0 bytes).")

    if file_size > max_bytes:
        raise ValidationException(
            f"File size of {file_size} bytes exceeds maximum allowed limit of {max_bytes} bytes.",
            details={"file_size": file_size, "max_bytes": max_bytes},
        )

    file_hash = hashlib.sha256(content).hexdigest()
    lower_filename = filename.lower()

    if lower_filename.endswith(".xlsm"):
        raise ValidationException(
            "Macro-enabled Excel workbooks (.xlsm) are strictly prohibited for security reasons.",
            details={"filename": filename},
        )

    if lower_filename.endswith(".xlsx"):
        sheet_names, selected_sheet, raw_headers, headers, rows = parse_xlsx_content(
            content=content,
            max_rows=max_rows,
            target_sheet=sheet_name,
        )
        return ParsedImport(
            file_hash=file_hash,
            file_size=file_size,
            headers=headers,
            raw_headers=raw_headers,
            rows=rows,
            sheet_names=sheet_names,
            selected_sheet=selected_sheet,
        )
    elif lower_filename.endswith((".csv", ".tsv", ".txt")):
        raw_headers, headers, rows = parse_csv_content(
            content=content,
            max_rows=max_rows,
        )
        return ParsedImport(
            file_hash=file_hash,
            file_size=file_size,
            headers=headers,
            raw_headers=raw_headers,
            rows=rows,
            sheet_names=[],
            selected_sheet=None,
        )
    else:
        raise ValidationException(
            "Unsupported file format. Only CSV (.csv) and Excel (.xlsx) files are supported.",
            details={"filename": filename},
        )
