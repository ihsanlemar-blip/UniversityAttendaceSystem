"""Test suite for safe spreadsheet exports, formula injection defense, and format parity."""

import io
import uuid

import openpyxl  # type: ignore[import-untyped]
import pytest
from fastapi.testclient import TestClient

from backend.app.models.university import University
from backend.app.models.user import User
from backend.app.reports.exports import sanitize_cell_value
from backend.tests.test_attendance_corrections import setup_closed_session
from backend.tests.test_attendance_policy import setup_class_occurrence
from backend.tests.test_people import create_test_user
from backend.tests.test_users import get_admin_headers


def test_sanitize_cell_value_unit() -> None:
    """Verify formula injection sanitizer prefixes quote on dangerous symbols."""
    assert sanitize_cell_value("=1+1") == "'=1+1"
    assert sanitize_cell_value("+2+2") == "'+2+2"
    assert sanitize_cell_value("-3-3") == "'-3-3"
    assert sanitize_cell_value("@SUM(A1:A10)") == "'@SUM(A1:A10)"
    assert sanitize_cell_value("\tTAB") == "'\tTAB"
    assert sanitize_cell_value("\rCR") == "'\rCR"
    assert sanitize_cell_value("CS-101") == "CS-101"
    assert sanitize_cell_value("Normal Student Name") == "Normal Student Name"
    assert sanitize_cell_value(123) == 123
    assert sanitize_cell_value(95.5) == 95.5
    assert sanitize_cell_value(None) == ""


@pytest.mark.asyncio
async def test_course_roster_csv_and_xlsx_export(
    client: TestClient,
    test_university: University,
    test_admin_user: User,
) -> None:
    """Verify CSV (UTF-8 BOM + formula defense) and XLSX exports for course roster."""
    admin_headers = get_admin_headers(client, test_admin_user, test_university)
    _, offering_id, occurrence_id = setup_class_occurrence(client, admin_headers)

    # Create a student with a formula prefix in number
    salt = uuid.uuid4().hex[:6]
    uid, _, _ = create_test_user(client, admin_headers, prefix=f"inj_{salt}")
    dangerous_student_num = f"=CMD_{salt}"
    stud_res = client.post(
        "/api/v1/students",
        headers=admin_headers,
        json={
            "user_id": uid,
            "student_number": dangerous_student_num,
            "admission_date": "2024-09-01",
        },
    )
    assert stud_res.status_code == 201
    s_id = stud_res.json()["data"]["id"]

    roles_res = client.get("/api/v1/roles", headers=admin_headers)
    student_role_id = next(r["id"] for r in roles_res.json()["data"] if r["code"] == "STUDENT")
    client.post(
        f"/api/v1/users/{uid}/role-assignments",
        headers=admin_headers,
        json={
            "role_id": student_role_id,
            "scope_type": "UNIVERSITY",
            "scope_id": str(test_university.id),
        },
    )
    client.post(
        f"/api/v1/course-offerings/{offering_id}/enrollments",
        headers=admin_headers,
        json={"student_id": s_id},
    )

    setup_closed_session(client, admin_headers, occurrence_id)

    # 1. Export CSV
    csv_res = client.get(
        f"/api/v1/reports/attendance/course-offerings/{offering_id}/export?format=CSV",
        headers=admin_headers,
    )
    assert csv_res.status_code == 200
    assert csv_res.headers["content-type"] == "text/csv; charset=utf-8"
    assert "attachment; filename=" in csv_res.headers["content-disposition"]

    csv_bytes = csv_res.content
    # Verify UTF-8 BOM
    assert csv_bytes.startswith(b"\xef\xbb\xbf")

    csv_text = csv_bytes.decode("utf-8-sig")
    # Verify formula injection defense escaped the '='
    assert f"'{dangerous_student_num}" in csv_text

    # 2. Export XLSX
    xlsx_res = client.get(
        f"/api/v1/reports/attendance/course-offerings/{offering_id}/export?format=XLSX",
        headers=admin_headers,
    )
    assert xlsx_res.status_code == 200
    assert "openxmlformats" in xlsx_res.headers["content-type"]

    xlsx_bytes = xlsx_res.content
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert wb.active.title == "Course Roster Report"


@pytest.mark.asyncio
async def test_session_and_department_exports(
    client: TestClient,
    test_university: University,
    test_admin_user: User,
) -> None:
    """Verify session and department export endpoints in both CSV and XLSX."""
    admin_headers = get_admin_headers(client, test_admin_user, test_university)
    _, offering_id, occurrence_id = setup_class_occurrence(client, admin_headers)
    session_id, _ = setup_closed_session(client, admin_headers, occurrence_id)

    # 1. Session CSV
    s_csv = client.get(
        f"/api/v1/reports/attendance/sessions/{session_id}/export?format=CSV",
        headers=admin_headers,
    )
    assert s_csv.status_code == 200
    assert s_csv.content.startswith(b"\xef\xbb\xbf")

    # 2. Session XLSX
    s_xlsx = client.get(
        f"/api/v1/reports/attendance/sessions/{session_id}/export?format=XLSX",
        headers=admin_headers,
    )
    assert s_xlsx.status_code == 200
    wb_sess = openpyxl.load_workbook(io.BytesIO(s_xlsx.content))
    assert wb_sess.active.title == "Session Report"

    # 3. Create Department and test department export
    dept_res = client.post(
        "/api/v1/academic-units",
        headers=admin_headers,
        json={
            "name": f"Dept of Biology {uuid.uuid4().hex[:4]}",
            "code": f"BIO_{uuid.uuid4().hex[:4]}".upper(),
            "unit_type": "DEPARTMENT",
        },
    )
    assert dept_res.status_code == 201
    dept_id = dept_res.json()["data"]["id"]

    d_csv = client.get(
        f"/api/v1/reports/attendance/department/{dept_id}/export?format=CSV",
        headers=admin_headers,
    )
    assert d_csv.status_code == 200
    assert d_csv.content.startswith(b"\xef\xbb\xbf")

    d_xlsx = client.get(
        f"/api/v1/reports/attendance/department/{dept_id}/export?format=XLSX",
        headers=admin_headers,
    )
    assert d_xlsx.status_code == 200
    wb_dept = openpyxl.load_workbook(io.BytesIO(d_xlsx.content))
    assert wb_dept.active.title == "Department Aggregate"
